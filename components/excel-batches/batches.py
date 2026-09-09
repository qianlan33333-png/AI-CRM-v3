"""Excel preparation and observation service. Never sends or approves messages."""
import argparse
import hashlib
import hmac
import io
import json
import os
import re
import sqlite3
import time
import threading
import urllib.parse
import zipfile
from datetime import datetime, timezone, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from contextlib import contextmanager

HEADERS = ['unionid', '话术', '小程序 path', '发送人 userid', '标题']
WINDOWS = (12, 24, 48)
MAX_BYTES = 8 * 1024 * 1024


def stamp():
    return datetime.now(timezone.utc).isoformat()


def instant(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(timezone.utc)


def digest(raw):
    return 'sha256:' + hashlib.sha256(raw).hexdigest()


class Invalid(ValueError):
    pass


def parse_excel(raw):
    from openpyxl import load_workbook
    if not raw or len(raw) > MAX_BYTES:
        raise Invalid('文件为空或超过 8 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 64 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise Invalid('Excel 解压后过大')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Invalid:
        raise
    except Exception:
        raise Invalid('请上传有效的 .xlsx 文件') from None
    try:
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > 5001:
            raise Invalid('每批最多 5000 行')
        rows = sheet.iter_rows()
        header = next(rows, ())
        if [c.value for c in header] != HEADERS:
            raise Invalid('首行必须依次为：' + '、'.join(HEADERS))
        result, seen, text_bytes = [], set(), 0
        for number, cells in enumerate(rows, 2):
            if not any(c.value is not None for c in cells):
                continue
            if len(cells) != 5 or any(c.data_type == 'f' or (not isinstance(c.value, str) and not (i == 4 and c.value is None)) for i, c in enumerate(cells)):
                raise Invalid(f'第 {number} 行：五列必须是文本，不能使用公式或数字格式的 ID')
            unionid, text, path, sender, title = [c.value for c in cells]
            title = (title or "").strip()
            unionid, path, sender = unionid.strip(), path.strip(), sender.strip()
            if not all((unionid, text.strip(), path, sender)) or any(re.search(r'[\x00-\x20]', x) for x in (unionid, path, sender)):
                raise Invalid(f'第 {number} 行：存在空值或无效标识')
            if len(unionid.encode()) > 256 or len(sender.encode()) > 256 or len(path.encode()) > 1024 or len(text.encode()) > 8000 or len(title.encode()) > 512:
                raise Invalid(f'第 {number} 行：内容超长')
            if '://' in path or path.startswith('//') or not path.startswith('pages/'):
                raise Invalid(f'第 {number} 行：需要完整的小程序 pages/ 路径')
            if unionid in seen:
                raise Invalid(f'第 {number} 行：同一批次的 UnionID 重复，请先合并为一行')
            text_bytes += sum(len(v.encode()) for v in (unionid, text, path, sender, title))
            if text_bytes > 8 * 1024 * 1024:
                raise Invalid('单批次五列文本总量不能超过 8 MB')
            seen.add(unionid)
            result.append({'unionid': unionid, 'text': text, 'path': path, 'sender_userid': sender, 'title': title})
            if len(result) > 5000:
                raise Invalid('每批最多 5000 行')
        if not result:
            raise Invalid('Excel 没有接收用户')
        return result
    finally:
        book.close()


def content_key(path):
    """Only known routes are measurable. Never treat a generic app visit as an open."""
    parsed = urllib.parse.urlsplit(path)
    query = urllib.parse.parse_qs(parsed.query)
    if parsed.path == 'pages/article/article' and len(query.get('lesson_id', [])) == 1:
        return ('lesson', query['lesson_id'][0])
    if parsed.path in ('pages/case/case', 'pages/case-detail/case-detail') and len(query.get('case_id', [])) == 1:
        return ('case', query['case_id'][0])
    return None


class Source:
    """Read-only HXC adapter; SQL is operator configuration, never HTTP input.

    Classification SQL reuses the existing A/B/C/D projection, rather than
    inventing new activity thresholds. Missing/ambiguous facts stay unknown.
    """
    def __init__(self, config):
        self.config = config
        self.local = threading.local()

    @contextmanager
    def session(self):
        self.local.batch = True
        self.local.connection = None
        try:
            yield
        finally:
            if self.local.connection is not None:
                self.local.connection.close()
            self.local.connection = None
            self.local.batch = False

    def query(self, name, params):
        import pymysql
        query = self.config.get(name)
        if not query or not query.lstrip().upper().startswith('SELECT '):
            raise RuntimeError('source_unconfigured')
        batch = getattr(self.local, 'batch', False)
        db = getattr(self.local, 'connection', None) if batch else None
        if db is None:
            db = pymysql.connect(**self.config['mysql'], cursorclass=pymysql.cursors.DictCursor,
                                 connect_timeout=5, read_timeout=20, write_timeout=5, autocommit=False)
            if batch:
                self.local.connection = db
        try:
            with db.cursor() as cursor:
                cursor.execute('SET TRANSACTION READ ONLY')
                cursor.execute('START TRANSACTION')
                cursor.execute(query, params)
                # Full result for this recipient/content/window; no global LIMIT.
                return cursor.fetchall()
        finally:
            db.rollback()
            if not batch:
                db.close()

    def segment(self, unionid):
        try:
            rows = self.query('segment_sql', (unionid,))
            if len(rows) == 1 and rows[0].get('segment') in 'ABCD' and len(rows[0]['segment']) == 1:
                return rows[0]['segment']
        except Exception:
            pass
        return 'unknown'

    def opens(self, unionid, path, start, end):
        key = content_key(path)
        if not key:
            return None
        try:
            # Empty event rows prove zero opens only when the collector confirms
            # complete coverage of this content and interval.
            coverage = self.query('coverage_sql', (key[0], key[1], start, end))
            if len(coverage) != 1 or coverage[0].get('complete') != 1:
                return None
            users = self.query('user_sql', (unionid,))
            if len(users) != 1:
                return None
            # Each configured query takes user_id, content_id, inclusive start/end.
            rows = self.query(key[0] + '_opens_sql', (users[0]['user_id'], key[1], start, end))
            return [instant(row['opened_at']).isoformat() for row in rows]
        except Exception:
            return None


class Service:
    def __init__(self, database, config, source=None):
        self.database, self.config = str(database), config
        self.source = source or Source(config.get('source', {}))
        with self.db() as db:
            db.executescript('''
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS imports (
              batch_key TEXT PRIMARY KEY, file_digest TEXT NOT NULL, created_at TEXT NOT NULL,
              rows_json TEXT NOT NULL, plan_id INTEGER UNIQUE, request_key TEXT UNIQUE);
            CREATE INDEX IF NOT EXISTS imports_digest ON imports(file_digest,created_at);
            CREATE TABLE IF NOT EXISTS covers (digest TEXT PRIMARY KEY, body BLOB NOT NULL, mime TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS snapshots (snapshot_key TEXT PRIMARY KEY, created_at TEXT NOT NULL, groups_json TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS observations (plan_id INTEGER PRIMARY KEY, body TEXT NOT NULL, updated_at TEXT NOT NULL);
            ''')

    @contextmanager
    def db(self):
        db = sqlite3.connect(self.database, timeout=30)
        db.execute('PRAGMA foreign_keys=ON')
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def cover(self, raw):
        if len(raw) > 2 * 1024 * 1024:
            raise Invalid('封面超过 2 MB')
        if raw.startswith(b'\x89PNG\r\n\x1a\n'):
            mime = 'image/png'
        elif raw.startswith(b'\xff\xd8\xff'):
            mime = 'image/jpeg'
        else:
            raise Invalid('封面必须是 PNG 或 JPEG')
        key = digest(raw)
        with self.db() as db:
            db.execute('INSERT OR IGNORE INTO covers VALUES(?,?,?)', (key, raw, mime))
        return key

    def imported(self, row, replayed):
        return {'batch_key': row['batch_key'], 'file_digest': row['file_digest'], 'created_at': row['created_at'],
                'plan_id': row['plan_id'], 'rows': json.loads(row['rows_json']), 'replayed': replayed}

    def import_file(self, raw, new_batch=False, request_key=''):
        file_digest = digest(raw)
        if new_batch and (not re.fullmatch(r'[A-Za-z0-9_-]{8,100}', request_key)):
            raise Invalid('新批次需要稳定的请求编号')
        with self.db() as db:
            prior = db.execute('SELECT * FROM imports WHERE request_key=?', (request_key or None,)).fetchone()
            if not prior and not new_batch:
                prior = db.execute('SELECT * FROM imports WHERE file_digest=? ORDER BY created_at LIMIT 1', (file_digest,)).fetchone()
            if prior:
                if prior['file_digest'] != file_digest:
                    raise Invalid('同一请求编号不能更换文件')
                return self.imported(prior, True)
        rows = parse_excel(raw)
        for row in rows:
            row['card'] = {'appid': self.config['appid'], 'path': row['path'],
                           'title': row['title'], 'cover_digest': ''}
        batch_key = hashlib.sha256((file_digest + (request_key if new_batch else '')).encode()).hexdigest()
        with self.db() as db:
            db.execute('INSERT OR IGNORE INTO imports(batch_key,file_digest,created_at,rows_json,request_key) VALUES(?,?,?,?,?)',
                       (batch_key, file_digest, stamp(), json.dumps(rows, ensure_ascii=False), request_key or None))
            result = db.execute('SELECT * FROM imports WHERE batch_key=?', (batch_key,)).fetchone()
            return self.imported(result, False)

    def link(self, batch_key, plan_id):
        with self.db() as db:
            changed = db.execute('UPDATE imports SET plan_id=? WHERE batch_key=? AND (plan_id IS NULL OR plan_id=?)',
                                 (plan_id, batch_key, plan_id)).rowcount
            if changed != 1:
                raise Invalid('批次关联冲突')
        return {'ok': True}

    def snapshot(self, key, rows):
        with self.db() as db:
            prior = db.execute('SELECT groups_json FROM snapshots WHERE snapshot_key=?', (key,)).fetchone()
            if prior:
                return {'snapshot_key': key}
        groups = {str(row['id']): self.source.segment(row['unionid']) for row in rows}
        with self.db() as db:
            db.execute('INSERT OR IGNORE INTO snapshots VALUES(?,?,?)', (key, stamp(), json.dumps(groups)))
        return {'snapshot_key': key}

    def observe(self, plan_id, snapshot_key, rows, now=None):
        now = instant(now) if now else datetime.now(timezone.utc)
        with self.db() as db:
            snapshot = db.execute('SELECT groups_json FROM snapshots WHERE snapshot_key=?', (snapshot_key,)).fetchone()
        groups = json.loads(snapshot[0]) if snapshot else {}
        counts, details = {}, []
        report = {str(h): {g: {'sent': 0, 'matured': 0, 'measurable': 0, 'opened': 0, 'observing': 0, 'unavailable': 0}
                          for g in ('A', 'B', 'C', 'D', 'unknown')} for h in WINDOWS}
        for row in rows:
            state = row['state']
            counts[state] = counts.get(state, 0) + 1
            group = groups.get(str(row['id']), 'unknown')
            item = dict(row, segment=group, windows={})
            details.append(item)
            if state != 'delivery_proven' or not row.get('sent_at'):
                continue
            sent = instant(row['sent_at'])
            # Re-read the bounded 48h interval on each pass to include late events.
            events = self.source.opens(row['unionid'], row['path'], sent, min(now, sent + timedelta(hours=48)))
            times = [instant(x) for x in events] if events is not None else None
            for h in WINDOWS:
                stats = report[str(h)][group]
                stats['sent'] += 1
                end = sent + timedelta(hours=h)
                if now < end:
                    stats['observing'] += 1
                    item['windows'][str(h)] = 'observing'
                    continue
                stats['matured'] += 1
                if times is None:
                    stats['unavailable'] += 1
                    item['windows'][str(h)] = 'unavailable'
                    continue
                stats['measurable'] += 1
                opened = any(sent <= t <= end for t in times)
                stats['opened'] += int(opened)
                item['windows'][str(h)] = 'opened' if opened else 'not_opened'
        for window in report.values():
            for stats in window.values():
                # A missing source must never silently shrink the agreed denominator.
                stats['open_rate'] = stats['opened'] / stats['matured'] if stats['matured'] and not stats['unavailable'] else None
        result = {'plan_id': plan_id, 'updated_at': now.isoformat(), 'counts': counts, 'windows': report, 'rows': details}
        with self.db() as db:
            db.execute('INSERT INTO observations VALUES(?,?,?) ON CONFLICT(plan_id) DO UPDATE SET body=excluded.body,updated_at=excluded.updated_at',
                       (plan_id, json.dumps(result, ensure_ascii=False), now.isoformat()))
        return result

    def report(self, plan_id):
        with self.db() as db:
            row = db.execute('SELECT body FROM observations WHERE plan_id=?', (plan_id,)).fetchone()
        return json.loads(row[0]) if row else {'plan_id': plan_id, 'pending': True, 'rows': [], 'windows': {}}


class Handler(BaseHTTPRequestHandler):
    server_version = 'ExcelBatch'
    def log_message(self, *args):
        pass  # Do not log identifiers, request bodies or credentials.

    def do_GET(self):
        self.handle_api()

    def do_POST(self):
        self.handle_api()

    def handle_api(self):
        with self.server.service.source.session():
            self._handle_api()

    def _handle_api(self):
        try:
            expected = 'Bearer ' + self.server.token
            if not hmac.compare_digest(self.headers.get('Authorization', ''), expected):
                return self.reply(401, {'error': 'unauthorized'})
            path = urllib.parse.urlsplit(self.path)
            service = self.server.service
            if self.command == 'GET' and path.path == '/health':
                return self.reply(200, {'ok': True})
            if self.command == 'GET' and path.path.startswith('/covers/'):
                key = urllib.parse.unquote(path.path[len('/covers/'):])
                with service.db() as db:
                    row = db.execute('SELECT body,mime FROM covers WHERE digest=?', (key,)).fetchone()
                if not row:
                    return self.reply(404, {'error': 'cover_not_found'})
                return self.reply(200, row['body'], row['mime'])
            if self.command == 'GET' and path.path.startswith('/reports/'):
                return self.reply(200, service.report(int(path.path.rsplit('/', 1)[1])))
            if self.command != 'POST':
                return self.reply(404, {'error': 'not_found'})
            length = int(self.headers.get('Content-Length', '0'))
            if length < 1 or length > 24 * 1024 * 1024:
                return self.reply(413, {'error': 'body_too_large'})
            body = self.rfile.read(length)
            if path.path == '/imports':
                query = urllib.parse.parse_qs(path.query)
                return self.reply(200, service.import_file(body, query.get('new') == ['1'], self.headers.get('Idempotency-Key', '')))
            if path.path == '/covers':
                return self.reply(200, {'cover_digest': service.cover(body)})
            data = json.loads(body)
            if path.path == '/link':
                return self.reply(200, service.link(data['batch_key'], int(data['plan_id'])))
            if path.path == '/snapshots':
                return self.reply(200, service.snapshot(data['snapshot_key'], data['rows']))
            if path.path == '/observations':
                return self.reply(200, service.observe(int(data['plan_id']), data['snapshot_key'], data['rows']))
            return self.reply(404, {'error': 'not_found'})
        except Invalid as error:
            return self.reply(400, {'error': 'invalid_input', 'message': str(error)})
        except (ValueError, KeyError):
            return self.reply(400, {'error': 'invalid_input'})
        except Exception:
            return self.reply(503, {'error': 'component_unavailable'})

    def reply(self, status, body, mime='application/json'):
        raw = body if isinstance(body, bytes) else json.dumps(body, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header('Content-Type', mime)
        self.send_header('Content-Length', str(len(raw)))
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(raw)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--config', required=True)
    parser.add_argument('--database', required=True)
    parser.add_argument('--port', type=int, default=8791)
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text())
    token = os.environ.get('EXCEL_BATCH_TOKEN', '')
    if len(token) < 32 or not config.get('appid'):
        raise SystemExit('Configure token and appid before starting')
    service = Service(args.database, config)
    server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
    server.service, server.token = service, token
    server.serve_forever()


if __name__ == '__main__':
    main()
